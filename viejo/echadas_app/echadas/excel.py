"""Exportación Excel estructurada (openpyxl) de una echada GENERADA."""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .dibujo import FILL, tirada_drawing, tirada_titulo
from .generador import Echada, pos_info
from .render import to_svg
from .tablas import ABBR, composicion, posiciones, resumen

HEAD_FILL = PatternFill("solid", fgColor="1F2937")
BAND_FILL = PatternFill("solid", fgColor="F1F3F4")
ACCENT = "B3261E"
RED = "E00000"
THIN = Side(style="thin", color="C9CCD1")
THICK = Side(style="medium", color="111111")


def _svg_png(svg: str, zoom: float = 2.0):
    try:
        import pymupdf
    except ImportError:
        try:
            import fitz as pymupdf
        except ImportError:
            return None
    try:
        doc = pymupdf.open(stream=svg.encode("utf-8"), filetype="svg")
        return doc[0].get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False).tobytes("png")
    except Exception:
        return None


def build_xlsx(ech: Echada, job, *, include_image=True, repite_8=True, control: str | None = None,
               diferencias: list | None = None) -> bytes:
    wb = Workbook()

    # ---------------------------------------------------------------- Orden
    ws = wb.active
    ws.title = "Orden"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 70
    ws["A1"] = "PLANO DE IMPOSICIÓN · ECHADA FULL COLOR"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"{ech.total} páginas  ·  cuerpos {ech.config}  ·  {len(ech.tiradas)} tirada(s)"
    ws["A2"].font = Font(bold=True, size=12, color=ACCENT)
    r = 4

    def section(title):
        nonlocal r
        for c in (1, 2):
            ws.cell(r, c).fill = HEAD_FILL
        ws.cell(r, 1, title).font = Font(bold=True, color="FFFFFF")
        r += 1

    def kv(k, v):
        nonlocal r
        a, b = ws.cell(r, 1, k), ws.cell(r, 2, v)
        a.font = Font(color="5F6368")
        b.font = Font(bold=True)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        a.border = b.border = Border(bottom=THIN)
        r += 1

    section("Datos del trabajo")
    for k, v in [("Publicación / Producto", job.publicacion), ("Edición / Fecha", job.edicion),
                 ("Orden de trabajo", job.orden), ("Responsable", job.responsable),
                 ("Destino / Sección", job.destino), ("Observaciones", job.observaciones)]:
        kv(k, v or "—")
    r += 1
    section("Resumen técnico (generado)")
    for k, v in resumen(ech):
        kv(k, v)
    for t in ech.tiradas:
        titulo, detalle = tirada_titulo(ech, t)
        kv(f"Tirada {t.numero}", detalle + ("  ·  REPITE" if repite_8 and t.formato == 8 else ""))
    kv("Generado", datetime.now().strftime("%d/%m/%Y %H:%M"))
    if control:
        kv("Control contra plantilla", control)
    for a in ech.avisos:
        kv("Aviso", a)
    r += 1
    section("Composición")
    hdr = ["#", "Cuerpo", "Figura", "Págs.", "Reparto en tiradas"]
    for row in composicion(ech):
        kv(f"{row['#']}. {row['Cuerpo']}", f"{row['Págs.']} págs. · {row['Figura']} · {row['Reparto en tiradas']}")

    # ------------------------------------------------------- Plano (celdas)
    wp = wb.create_sheet("Plano")
    wp.sheet_view.showGridLines = False
    # columnas: por cara -> [marca, p, p, sep, marca, p, p]  + separador entre caras
    widths = []
    layout_cols = []  # (cara, col_imp or None, kind)
    for cara in ("L10", "L13"):
        cols = [(0, 1), (2, 3)] if cara == "L10" else [(3, 2), (1, 0)]
        for qi, pair in enumerate(cols):
            layout_cols.append((cara, pair, "mark"))
            layout_cols.append((cara, pair[0], "page"))
            layout_cols.append((cara, pair[1], "page"))
            layout_cols.append((cara, None, "gap"))
        layout_cols.append((None, None, "facegap"))
    for i, (cara, c, kind) in enumerate(layout_cols, start=1):
        wp.column_dimensions[get_column_letter(i)].width = {"mark": 4, "page": 8, "gap": 2, "facegap": 4}[kind]
    row = 1
    wp.cell(row, 1, f"Echada {ech.config} = {ech.total} págs.").font = Font(bold=True, size=14)
    row += 2
    for tir in reversed(ech.tiradas):
        titulo, detalle = tirada_titulo(ech, tir)
        c = wp.cell(row, 1, titulo + ("  ·  REPITE" if repite_8 and tir.formato == 8 else ""))
        c.font = Font(bold=True, size=12, color=ACCENT)
        wp.cell(row + 1, 1, detalle).font = Font(italic=True, size=9, color="5F6368")
        row += 3
        for web in tir.webs:
            # encabezados
            for i, (cara, c_, kind) in enumerate(layout_cols, start=1):
                if kind == "mark":
                    pair = c_
                    visible = (not tir.medio_formato) or 2 in pair
                    if visible:
                        h, _ = pos_info(cara, pair[0])
                        cell = wp.cell(row, i + 1, h)
                        cell.font = Font(bold=True, color=RED)
                        cell.alignment = Alignment(horizontal="center")
                        wp.merge_cells(start_row=row, start_column=i + 1, end_row=row, end_column=i + 2)
                        m = wp.cell(row + 1, i, web)
                        m.font = Font(bold=True, color=RED, size=13)
                        m.alignment = Alignment(horizontal="center", vertical="center")
                        wp.merge_cells(start_row=row + 1, start_column=i, end_row=row + 2, end_column=i)
                if kind == "page":
                    visible = (not tir.medio_formato) or c_ in (2, 3)
                    if not visible:
                        continue
                    left_of_pair = (cara == "L10" and c_ in (0, 2)) or (cara == "L13" and c_ in (3, 1))
                    for k, fila in enumerate(("top", "bot")):
                        p = ech.en(tir.numero, web, cara, c_, fila)
                        rep = False
                        if p is None and repite_8 and tir.formato == 8 and c_ in (0, 1):
                            p = ech.en(tir.numero, web, cara, c_ + 2, fila)
                            rep = p is not None
                        cell = wp.cell(row + 1 + k, i)
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.border = Border(right=THICK if left_of_pair else None,
                                             bottom=THICK if fila == "top" else None)
                        if p is None:
                            continue
                        cu = ech.cuerpos[p.cuerpo]
                        cell.value = p.pagina
                        cell.font = Font(bold=True, size=14, color="9CA3AF" if rep else "111111")
                        if cu.figura in FILL:
                            cell.fill = PatternFill("solid", fgColor=FILL[cu.figura].lstrip("#").replace(
                                "FFFFFF", "EDEDED"))
                            cell.value = f"{ABBR.get(cu.figura, '')}{p.pagina}"
            wp.row_dimensions[row + 1].height = 26
            wp.row_dimensions[row + 2].height = 26
            # LADO
            for cara, start in (("L10", 1), ("L13", 1 + 9)):
                cell = wp.cell(row + 3, start + 3, "LADO " + cara[1:])
                cell.font = Font(bold=True, color=RED)
            row += 6
        row += 1
    # leyenda
    row += 1
    wp.cell(row, 1, "Referencias:").font = Font(bold=True)
    row += 1
    for cu in ech.cuerpos:
        cell = wp.cell(row, 2, f"{ABBR.get(cu.figura, '')}n" if cu.figura in ABBR else "n")
        cell.alignment = Alignment(horizontal="center")
        if cu.figura in FILL:
            cell.fill = PatternFill("solid", fgColor=FILL[cu.figura].lstrip("#").replace("FFFFFF", "EDEDED"))
        wp.cell(row, 3, f"{cu.nombre} — {cu.figura} — {cu.paginas} págs.")
        row += 1

    # ---------------------------------------------------- Plano (imagen)
    if include_image:
        wi = wb.create_sheet("Plano (imagen)")
        wi.sheet_view.showGridLines = False
        wi["A1"] = f"Plano visual — {ech.config} = {ech.total} págs."
        wi["A1"].font = Font(bold=True, size=13)
        rr, ok = 3, False
        try:
            from openpyxl.drawing.image import Image as XLImage
            for tir in reversed(ech.tiradas):
                d = tirada_drawing(ech, tir, repite_8)
                png = _svg_png(to_svg(d))
                if not png:
                    continue
                wi.cell(rr, 1, tirada_titulo(ech, tir)[0]).font = Font(bold=True, color=ACCENT)
                img = XLImage(io.BytesIO(png))
                img.width, img.height = d.width, d.height
                wi.add_image(img, f"A{rr + 1}")
                rr += int(d.height / 20) + 4
                ok = True
        except Exception:
            ok = False
        if not ok:
            wi["A3"] = "Instale 'pymupdf' y 'pillow' para incluir la imagen. Ver hoja 'Plano'."

    # ------------------------------------------------------- Posiciones
    wd = wb.create_sheet("Posiciones")
    filas = posiciones(ech)
    headers = list(filas[0].keys()) if filas else []
    wd.append(headers)
    for f in filas:
        wd.append([f[h] for h in headers])
    for i, h in enumerate(headers, 1):
        wd.column_dimensions[get_column_letter(i)].width = max(9, len(h) + 3)
        c = wd.cell(1, i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center")
    if filas:
        tbl = Table(displayName="Posiciones", ref=f"A1:{get_column_letter(len(headers))}{len(filas) + 1}")
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        wd.add_table(tbl)
    wd.freeze_panes = "A2"

    if diferencias is not None:
        wc = wb.create_sheet("Control")
        wc["A1"] = control or "Control contra plantilla"
        wc["A1"].font = Font(bold=True, size=12)
        if diferencias:
            hdr = list(diferencias[0].keys())
            for i, h in enumerate(hdr, 1):
                c = wc.cell(3, i, h)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = HEAD_FILL
                wc.column_dimensions[get_column_letter(i)].width = 22
            for j, d in enumerate(diferencias, start=4):
                for i, h in enumerate(hdr, 1):
                    wc.cell(j, i, d[h])
        else:
            wc["A3"] = "Sin diferencias: todas las posiciones coinciden con la plantilla."

    for w in wb.worksheets:
        w.page_setup.orientation = "landscape"
        w.sheet_properties.pageSetUpPr.fitToPage = True
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
